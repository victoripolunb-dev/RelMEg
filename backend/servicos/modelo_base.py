"""
Modelo Base Dinâmico — gabarito estrutural das extrações do TSE.

Carrega programaticamente a estrutura, a ordem exata das colunas, as larguras
e os estilos de cabeçalho do arquivo CONTRATO de coleta — o
"Modelo base de coleta - Parlamentares.xlsx" criado pelo operador
(este repositório mantém a cópia interna em backend/templates/MODELO BASE).
Todas as saídas geradas pelas rotas do RelMeg seguem este gabarito com
fidelidade visual.

O CONTRATO (fonte de verdade) é o arquivo do operador; a cópia interna cobre
clones/deploy sem a pasta externa. Original nunca é sobrescrito.

Garantias:
    - O arquivo-fonte é apenas LEITURA (nunca é sobrescrito).
    - Ordem das colunas é a do contrato, com os cabeçalhos multi-linha (FPE,
      FCS, FPBio, FPEvang) e os estilos replicados na saída.
    - Se ambos os arquivos forem indisponíveis em deploy, cai num gabarito de
      fallback com as mesmas 25 colunas na mesma ordem — a extração nunca
      quebra por falta do modelo.
"""
from __future__ import annotations

import io
from pathlib import Path
from typing import Any, Dict, Optional

import openpyxl
import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from config import settings

# CONTRATO canônico: o arquivo criado pelo operador. Prioridade sobre a cópia
# interna (config.modelo_base_template) porque é a versão validada por ele.
NOME_ARQUIVO_CONTRATO = "Modelo base de coleta - Parlamentares.xlsx"
MODELO_BASE_OPERADOR = settings.repo_root / "modelo base" / NOME_ARQUIVO_CONTRATO

# ---------------------------------------------------------------------------
# Fallback (usado apenas se o arquivo não estiver acessível)
# ---------------------------------------------------------------------------

# Cabeçalhos multi-linha idênticos ao modelo (quebra de linha real dentro da
# célula, renderizada com wrap_text no Excel).
COLUNAS_MODELO_BASE = [
    "Casa Legislativa",
    "Nome",
    "Partido",
    "UF",
    "Eleição/ Reeleição",
    "Membro da FPE \n2023 a 2026\n(Somente para reeleição)",
    "Membro da FCS \n2023 a 2026\n(Somente para reeleição)",
    "Membro da FPBio \n2023 a 2026\n(Somente para reeleição)",
    "Membro da FPEvangê\n2023 a 2026\n(Somente para reeleição)",
    "Sinergia FPE",
    "Eixo de Atuação",
    "Celular Parlamentar",
    "Assessoria",
    "Contato",
    "E-mail",
    "Gabinete/Endereço",
    "Rede Social",
    "Profissão",
    "Data de Nascimento",
    "Idade",
    "Cor/Raça",
    "Gênero",
    "Escolaridade",
    "Estado Civil ",
    "Observação",
]

# Larguras (em caracteres) espelhadas do modelo — aplicadas na saída.
_LARGURAS_FALLBACK = {
    0: 17.63, 1: 36.5, 2: 16.88, 3: 10.88, 4: 23.88, 5: 30.0, 6: 26.25,
    7: 26.38, 8: 28.75, 9: 20.88, 10: 22.0, 11: 22.5, 12: 16.5, 13: 17.25,
    14: 30.25, 15: 22.13, 16: 25.13, 17: 31.5, 18: 25.38, 19: 12.13,
    20: 22.38, 21: 14.0, 22: 18.0, 23: 14.0, 24: 62.13,
}


def contrato_modelo() -> Path:
    """Caminho do CONTRATO de coleta (arquivo do operador → cópia interna).

    Prioridade:
        1. ``modelo base/Modelo base de coleta - Parlamentares.xlsx`` (original);
        2. ``backend/templates/MODELO BASE`` (cópia interna versionada).

    Se nenhum existir, devolve a cópia interna (o chamador trata a ausência
    com ``contrato_disponivel()`` — o gabarito canônico é o fallback final).
    """
    if MODELO_BASE_OPERADOR.exists():
        return MODELO_BASE_OPERADOR
    return settings.modelo_base_template


def contrato_disponivel() -> bool:
    """True se o CONTRATO (operador ou cópia interna) estiver acessível."""
    try:
        return contrato_modelo().exists()
    except OSError:
        return False


def _caminho_modelo() -> Path:
    """Atalho interno (mantém única fonte: o contrato canônico)."""
    return contrato_modelo()


def _cadeia_rgb(cor: Any) -> Optional[str]:
    """Extrai o RGB hex do color openpyxl (robusto a None/theme)."""
    if cor is None:
        return None
    try:
        valor = getattr(cor, "rgb", None)
        if valor:
            return str(valor)
    except Exception:
        return None
    return None


def _ler_gabarito() -> Optional[Dict[str, Any]]:
    """Lê o arquivo MODELO BASE e devolve o gabarito, ou None se inacessível."""
    caminho = _caminho_modelo()
    try:
        if not caminho.exists():
            return None
        # BytesIO ignora a ausência de extensão (o arquivo não é .xlsx).
        wb = openpyxl.load_workbook(io.BytesIO(caminho.read_bytes()))
        sh = wb["Câmara"] if "Câmara" in wb.sheetnames else wb[wb.sheetnames[0]]

        colunas: list = []
        larguras: Dict[int, float] = {}
        estilos: Dict[int, Dict[str, Any]] = {}
        for col in range(1, sh.max_column + 1):
            cel = sh.cell(row=1, column=col)
            if cel.value is None:
                break
            colunas.append(str(cel.value))

            letra = get_column_letter(col)
            dim = sh.column_dimensions.get(letra)
            larguras[col - 1] = round(dim.width, 2) if dim and dim.width else 12.0

            fill = None
            if cel.fill and cel.fill.patternType:
                fill = _cadeia_rgb(cel.fill.fgColor)
            estilos[col - 1] = {
                "bold": bool(cel.font.bold),
                "fonte_cor": _cadeia_rgb(cel.font.color) or "FFFFFFFF",
                "fill": fill,
            }

        if not colunas or len(colunas) < 5:
            return None

        return {
            "origem": str(caminho),
            "colunas": colunas,
            "larguras": larguras,
            "estilos": estilos,
            "altura_cabecalho": round(sh.row_dimensions[1].height or 45.0, 1),
        }
    except Exception:
        return None


def _gabarito_fallback() -> Dict[str, Any]:
    """Gabarito garantido quando o MODELO BASE não está disponível."""
    return {
        "origem": "fallback (MODELO BASE indisponível)",
        "colunas": list(COLUNAS_MODELO_BASE),
        "larguras": dict(_LARGURAS_FALLBACK),
        "estilos": {i: {"bold": True, "fonte_cor": "FFFFFFFF", "fill": "FF1F3864"}
                    for i in range(len(COLUNAS_MODELO_BASE))},
        "altura_cabecalho": 45.0,
    }


_CACHE_GABARITO: Dict[str, Any] = {"mtime": None, "gabarito": None}


def gabarito_modelo() -> Dict[str, Any]:
    """Gabarito canônico (cacheado por mtime) — sempre o modelo vigente."""
    caminho = _caminho_modelo()
    try:
        mtime = caminho.stat().st_mtime if caminho.exists() else None
    except OSError:
        mtime = None

    cache = _CACHE_GABARITO
    if cache["mtime"] != mtime:
        cache["mtime"] = mtime
        cache["gabarito"] = _ler_gabarito()

    gabarito = cache["gabarito"]
    return gabarito if gabarito else _gabarito_fallback()


# ---------------------------------------------------------------------------
# Aplicação do gabarito na planilha de saída
# ---------------------------------------------------------------------------

# Índices (1-based nas planilhas) das colunas com formatação de dados especial.
_COL_DATA_NASCIMENTO = 19   # Data de Nascimento
_COL_IDADE = 20             # Idade


def aplicar_gabarito(planilha, gabarito: Optional[Dict[str, Any]] = None) -> None:
    """Aplica larguras, estilo de cabeçalho e formatos numéricos do modelo."""
    gab = gabarito or gabarito_modelo()
    total = len(gab["colunas"])

    for indice, largura in gab["larguras"].items():
        try:
            planilha.column_dimensions[get_column_letter(indice + 1)].width = largura
        except Exception:
            pass

    alinhamento = Alignment(vertical="center", horizontal="center", wrap_text=True)
    for col in range(1, total + 1):
        cel = planilha.cell(row=1, column=col)
        est = gab["estilos"].get(col - 1, {})
        try:
            cel.font = Font(bold=bool(est.get("bold", True)),
                            color=est.get("fonte_cor") or "FFFFFFFF")
            fill_hex = est.get("fill")
            if fill_hex:
                cel.fill = PatternFill(fill_type="solid", fgColor=str(fill_hex))
            cel.alignment = alinhamento
        except Exception:
            pass

    planilha.row_dimensions[1].height = gab["altura_cabecalho"]
    planilha.freeze_panes = "C2"

    # Formatos de dados: datas em dd/MM/yyyy e idade como inteiro.
    for linha in range(2, planilha.max_row + 1):
        if total >= _COL_DATA_NASCIMENTO:
            planilha.cell(row=linha, column=_COL_DATA_NASCIMENTO).number_format = "dd/MM/yyyy"
        if total >= _COL_IDADE:
            planilha.cell(row=linha, column=_COL_IDADE).number_format = "#,##0"


def gravar_com_gabarito(
    df: pd.DataFrame,
    caminho: Path,
    gabarito: Optional[Dict[str, Any]] = None,
) -> None:
    """Grava o .xlsx final seguindo fielmente o MODELO BASE."""
    df.to_excel(caminho, index=False, sheet_name="Candidatos")
    wb = openpyxl.load_workbook(str(caminho))
    aplicar_gabarito(wb.active, gabarito or gabarito_modelo())
    wb.save(str(caminho))


def bytes_com_gabarito(
    df: pd.DataFrame,
    gabarito: Optional[Dict[str, Any]] = None,
) -> bytes:
    """Serializa o .xlsx em memória (download) já aplicando o gabarito."""
    buffer = io.BytesIO()
    gab = gabarito or gabarito_modelo()
    with pd.ExcelWriter(buffer, engine="openpyxl") as escritor:
        df.to_excel(escritor, index=False, sheet_name="Candidatos")
        aplicar_gabarito(escritor.book["Candidatos"], gab)
    return buffer.getvalue()