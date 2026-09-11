"""
Exportador de Ficha Legislativa (.docx) — relmeg_core.

Gera, sob demanda, um relatório Word de uma proposição coletada no repositório,
reutilizando a identidade visual da casa (AGENTS.md → "Entregas ao cliente"):

    - Fonte Montserrat, texto em ``#333333``;
    - Número da proposição e links de destaque em negrito vermelho ``#ff0000``;
    - Gravado em ``~/Desktop/RelMeg - Entregas/Relatórios/`` (settings.dir_relatorios);
    - Nunca sobrescreve o "MODELO A SER SEGUIDO.docx".

O conteúdo é montado do ZERO (sem clonar o template): a ficha legislativa não é
um clipping semanal, e sim um extrato técnico — mas segue o mesmo contrato de
identidade (fonte/cores) descrito nas convenções do projeto.

Nenhuma API externa é consultada aqui: a rota de exportação lê apenas o
repositório local (alimentado sob demanda pelas rotas de coleta).
"""
from __future__ import annotations

import datetime as _dt
import re
from pathlib import Path
from typing import Any, Dict, List, Optional

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Pt, RGBColor

from config import settings

FONTE = "Montserrat"
TAMANHO_PT = 10
COR_TEXTO = RGBColor(0x33, 0x33, 0x33)  # 333333
COR_NUMERO = RGBColor(0xFF, 0x00, 0x00)  # ff0000

URL_CAMARA_HYPERLINK = "https://www.camara.leg.br/proposicoesWeb/fichadetramitacao?idProposicao={id}"
URL_SENADO_HYPERLINK = "https://www25.senado.leg.br/web/atividade/materias/-/materia/{codigo}"


class FichaLegislativaError(Exception):
    """Erro de negócio com mensagem amigável."""


def _dados(modelo: Any) -> Dict[str, Any]:
    """Normaliza modelo Pydantic OU dict de linha do repositório."""
    dump = getattr(modelo, "model_dump", None)
    if callable(dump):
        return dump(mode="json")
    return dict(modelo)


def _val(dados: Dict[str, Any], chave: str) -> Any:
    return dados.get(chave)


def _txt(valor: Any) -> str:
    if valor is None:
        return ""
    texto = str(valor).strip()
    return "—" if not texto else texto


def _data_br(valor: Any) -> str:
    """Converte data/ISO para DD/MM/AAAA (ou devolve '—')."""
    if not valor:
        return "—"
    texto = str(valor).split("T")[0]
    if isinstance(valor, _dt.date):
        return valor.strftime("%d/%m/%Y")
    encontro = re.match(r"^(\d{4})-(\d{2})-(\d{2})$", texto)
    if encontro:
        return f"{encontro.group(3)}/{encontro.group(2)}/{encontro.group(1)}"
    return texto or "—"


def _sanitizar(nome: str) -> str:
    """Remove caracteres inválidos para nome de arquivo no Windows."""
    return re.sub(r'[\\/:*?"<>|]', "_", nome).replace(" ", "_")


def _link_publico(projeto: Dict[str, Any]) -> Optional[str]:
    """URL pública de referência da proposição, conforme a fonte."""
    fonte = (projeto.get("fonte") or "").strip().lower()
    id_externo = str(projeto.get("id_externo") or "")
    if fonte == "camara":
        return URL_CAMARA_HYPERLINK.format(id=id_externo)
    if fonte == "senado":
        return URL_SENADO_HYPERLINK.format(codigo=id_externo)
    return projeto.get("url_origem")


# ---------------------------------------------------------------------------
# Primitivas de formatação (identidade da casa)
# ---------------------------------------------------------------------------


def _novo_run(par, texto: str, *, bold: bool = False, color: RGBColor = COR_TEXTO,
              size: int = TAMANHO_PT, italic: bool = False) -> Any:
    run = par.add_run(texto)
    run.font.name = FONTE
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.underline = False
    run.bold = bold
    run.italic = italic
    return run


def _novo_paragrafo(doc: Document):
    par = doc.add_paragraph()
    par.paragraph_format.space_after = Pt(2)
    return par


def _titulo(doc: Document, texto: str, size: int) -> None:
    par = doc.add_paragraph()
    par.paragraph_format.space_before = Pt(6)
    par.paragraph_format.space_after = Pt(6)
    _novo_run(par, texto, bold=True, size=size)


def _linha(doc: Document, label: str, valor: Any, *, url: Optional[str] = None,
           indentar: bool = False) -> None:
    """'Rótulo: valor'. Se ``url``, o valor vira hyperlink vermelho em negrito."""
    par = _novo_paragrafo(doc)
    if indentar:
        par.paragraph_format.left_indent = Pt(18)
    if label:
        _novo_run(par, f"{label}: ", bold=True)
    if url and valor:
        _hyperlink(par, _txt(valor), url)
    else:
        _novo_run(par, _txt(valor))


def _hyperlink(par, texto: str, url: str) -> None:
    """Adiciona um <w:hyperlink> com o texto em negrito vermelho (estilo da casa)."""
    r_id = par.part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)

    run = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    fontes = OxmlElement("w:rFonts")
    for attr in ("w:ascii", "w:cs", "w:eastAsia", "w:hAnsi"):
        fontes.set(qn(attr), FONTE)
    rPr.append(fontes)
    bold = OxmlElement("w:b"); bold.set(qn("w:val"), "1"); rPr.append(bold)
    cor = OxmlElement("w:color"); cor.set(qn("w:val"), "ff0000"); rPr.append(cor)
    sz = OxmlElement("w:sz"); sz.set(qn("w:val"), str(TAMANHO_PT * 2)); rPr.append(sz)
    run.append(rPr)
    t = OxmlElement("w:t")
    t.set(qn("xml:space"), "preserve")
    t.text = texto
    run.append(t)
    hyperlink.append(run)
    par._p.append(hyperlink)


# ---------------------------------------------------------------------------
# Montagem do documento
# ---------------------------------------------------------------------------


def _numero_proposicao(projeto: Dict[str, Any]) -> str:
    sigla = str(projeto.get("sigla_tipo") or "PL").upper()
    numero = str(projeto.get("numero") or "")
    ano = str(projeto.get("ano") or "")
    if numero and ano:
        return f"{sigla} {numero}/{ano}"
    return sigla or projeto.get("id_externo") or "—"


def _montar(doc: Document, projeto: Dict[str, Any],
            tramitacoes: List[Dict[str, Any]]) -> None:
    """Preenche o documento com os dados normalizados do repositório."""
    link = _link_publico(projeto)

    par = doc.add_paragraph()
    par.paragraph_format.space_after = Pt(2)
    _novo_run(par, "FICHA LEGISLATIVA", bold=True, size=14)

    linha = doc.add_paragraph()
    linha.paragraph_format.space_after = Pt(10)
    _novo_run(linha, "RelMeg — extração legislativa sob demanda", italic=True, size=9)

    _titulo(doc, "Identificação", 12)
    if link:
        _linha(doc, "Proposição", _numero_proposicao(projeto), url=link)
    else:
        _linha(doc, "Proposição", _numero_proposicao(projeto))
    _linha(doc, "Casa de origem", _txt(projeto.get("orgao_origem")))
    _linha(doc, "Fonte", _txt(projeto.get("fonte")))
    _linha(doc, "Data de apresentação", _data_br(projeto.get("data_apresentacao")))
    _linha(doc, "Ementa", _txt(projeto.get("ementa")))
    _linha(doc, "Autor(es)", _txt(projeto.get("autor_principal")))
    integrantes = _val(projeto, "integrantes")
    if integrantes:
        _linha(doc, "Coautores", _txt(", ".join(integrantes) if isinstance(integrantes, list) else integrantes))
    _linha(doc, "Tema/pauta", _txt(projeto.get("pauta_tematica")))
    _linha(doc, "Situação", _txt(projeto.get("situacao")))
    _linha(doc, "Comissão atual", _txt(projeto.get("comissao_atual")))
    _linha(doc, "Relator", _txt(projeto.get("relator")))
    if projeto.get("url_origem"):
        _linha(doc, "Fonte dos dados", _txt(projeto.get("url_origem")), url=projeto.get("url_origem"))

    _titulo(doc, "Tramitação", 12)
    eventos = tramitacoes or []
    if not eventos:
        par = _novo_paragrafo(doc)
        _novo_run(par, "Histórico de tramitação indisponível para esta fonte.", italic=True)
    for ev in eventos:
        _linha(
            doc,
            _data_br(ev.get("data_evento")),
            _txt(ev.get("descricao_fase") or ev.get("status")),
            url=ev.get("url_evento") or None,
        )
        detalhe = " — ".join(
            parte for parte in (
                _txt(ev.get("orgao_local")),
                _txt(ev.get("status")),
                _txt(ev.get("despacho")),
            ) if parte not in ("—", "")
        )
        if detalhe:
            par = _novo_paragrafo(doc)
            par.paragraph_format.left_indent = Pt(18)
            _novo_run(par, detalhe, size=9)

    rodape = doc.add_paragraph()
    rodape.paragraph_format.space_before = Pt(14)
    _novo_run(
        rodape,
        f"Gerado em {_dt.datetime.now().strftime('%d/%m/%Y %H:%M')} pelo motor relmeg_core — "
        "sob demanda (AGENTS.md).",
        size=8,
        italic=True,
    )


def _montar_parlamentar(doc: Document, dados: Dict[str, Any]) -> None:
    """Preenche a ficha com dados normalizados de um parlamentar."""

    def _link(valor: Any, campo: str) -> None:
        link = dados.get(campo)
        if link:
            _linha(doc, campo.replace("_", " ").capitalize(), _txt(link), url=link)

    par = doc.add_paragraph()
    par.paragraph_format.space_after = Pt(2)
    _novo_run(par, "FICHA DE PARLAMENTAR", bold=True, size=14)

    linha = doc.add_paragraph()
    linha.paragraph_format.space_after = Pt(10)
    _novo_run(linha, "RelMeg — extração legislativa sob demanda", italic=True, size=9)

    _titulo(doc, "Identificação", 12)
    _linha(doc, "Nome completo", _txt(dados.get("nome_completo")))
    _linha(doc, "Nome de urna", _txt(dados.get("nome_urna")))
    cargo = _txt(dados.get("cargo"))
    if cargo != "—":
        _linha(doc, "Cargo", cargo)
    _linha(doc, "Partido", _txt(dados.get("partido")))
    _linha(doc, "UF", _txt(dados.get("uf")))
    status = dados.get("status_ativo")
    _linha(doc, "Situação do mandato", "Ativo" if status else "Inativo")
    _linha(doc, "Início do mandato", _data_br(dados.get("mandato_inicio")))
    _linha(doc, "Fim do mandato", _data_br(dados.get("mandato_fim")))
    _linha(doc, "E-mail institucional", _txt(dados.get("email")))
    _linha(doc, "Fonte", _txt(dados.get("fonte")))
    if dados.get("url_origem"):
        _linha(doc, "Fonte dos dados", _txt(dados.get("url_origem")), url=dados.get("url_origem"))

    _titulo(doc, "Perfil oficial", 12)
    _link(dados.get("url_perfil"), "url_perfil")

    rodape = doc.add_paragraph()
    rodape.paragraph_format.space_before = Pt(14)
    _novo_run(
        rodape,
        f"Gerado em {_dt.datetime.now().strftime('%d/%m/%Y %H:%M')} pelo motor relmeg_core — "
        "sob demanda (AGENTS.md).",
        size=8,
        italic=True,
    )


def gerar_ficha_parlamentar(parlamentar: Any) -> Dict[str, Any]:
    """Gera e grava a Ficha de Parlamentar (.docx) do repositório local.

    Aceita ``ParlamentarModel`` ou linha de ``buscar_parlamentar``. Grava em
    ``settings.dir_relatorios`` — formato padrão da casa (Montserrat/333333).
    """
    dados = _dados(parlamentar)

    pasta = settings.dir_relatorios
    try:
        pasta.mkdir(parents=True, exist_ok=True)
    except OSError as exc:
        raise FichaLegislativaError(
            f"Não foi possível criar a pasta de entregas: {pasta}"
        ) from exc

    nome = str(dados.get("nome_completo") or "Parlamentar").strip().upper()
    nome = re.sub(r"[^A-Z0-9À-ÜÇÃÕÉÁÍÓÚÊÂÔ]+", "_", nome).strip("_") or "Parlamentar"
    fonte = str(dados.get("fonte") or "fonte").lower()
    data = _dt.datetime.now().strftime("%Y%m%d-%H%M%S")
    arquivo = f"Ficha_Parlamentar_{nome}_{fonte}_{data}.docx"
    destino = pasta / arquivo

    doc = Document()
    doc.core_properties.title = f"Ficha de Parlamentar — {dados.get('nome_completo')}"
    _montar_parlamentar(doc, dados)
    try:
        doc.save(str(destino))
    except OSError as exc:
        raise FichaLegislativaError(f"Falha ao gravar o relatório: {destino}") from exc

    return {
        "arquivo": arquivo,
        "caminho": str(destino),
        "gerado_em": _dt.datetime.now().isoformat(timespec="seconds"),
    }


def gerar_ficha(projeto: Any, tramitacoes: Any, *, nome_extra: str = "") -> Dict[str, Any]:
    """Gera e grava a Ficha Legislativa .docx de uma proposição do repositório.

    Aceita modelos Pydantic ou dicionários de linha (buscar_projeto/
    buscar_tramitacoes). Grava em ``settings.dir_relatorios`` e NUNCA sobrescreve
    o modelo do Clipping.

    Retorna metadados: arquivo, caminho e total de tramitações.
    """
    dados_projeto = _dados(projeto)
    eventos = [_dados(t) for t in (tramitacoes or [])]

    pasta = settings.dir_relatorios
    try:
        pasta.mkdir(parents=True, exist_ok=True)
    except OSError as exc:
        raise FichaLegislativaError(
            f"Não foi possível criar a pasta de entregas: {pasta}"
        ) from exc

    sigla = str(dados_projeto.get("sigla_tipo") or "PL").upper()
    numero = str(dados_projeto.get("numero") or 0)
    ano = str(dados_projeto.get("ano") or 0)
    fonte = str(dados_projeto.get("fonte") or "proposicao").lower()
    data = _dt.datetime.now().strftime("%Y%m%d-%H%M%S")
    extra = f"_{_sanitizar(nome_extra)}" if nome_extra else ""
    nome = f"Ficha_Legislativa_{sigla}_{numero}_{ano}_{fonte}_{data}{extra}.docx"
    destino = pasta / nome

    doc = Document()
    doc.core_properties.title = f"Ficha Legislativa {sigla} {numero}/{ano}"
    _montar(doc, dados_projeto, eventos)
    try:
        doc.save(str(destino))
    except OSError as exc:
        raise FichaLegislativaError(f"Falha ao gravar o relatório: {destino}") from exc

    return {
        "arquivo": nome,
        "caminho": str(destino),
        "total_tramitacoes": len(eventos),
        "gerado_em": _dt.datetime.now().isoformat(timespec="seconds"),
    }