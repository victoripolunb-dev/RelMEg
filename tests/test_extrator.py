"""Testes do extrator TSE com o Modelo Base Dinâmico (backend/extrator_tse.py)."""
import asyncio
import time

import httpx
import pandas as pd
import pytest

import database
from config import settings
from extrator_tse import (
    ExtrairTSEError,
    _atraso_tentativa,
    _bens_individuais,
    _bytes_xlsx,
    _eixo_atividade,
    _get_json,
    _idade,
    _id_eleicao,
    _norm_detalhe_rico,
    _normalizar_campos_detalhe,
    _validar_cargo,
    _validar_uf,
    _valor_data,
    chave_cache,
    estrutura_dataframe,
    obter_candidatos_cacheados,
)
from modelo_base import COLUNAS_MODELO_BASE, gabarito_modelo

CHAVE_GO = "tse|2026|GO|7|completo"


def _candidato(n: int = 1) -> dict:
    return {
        "id_candidato": n,
        "nome_urna": "NOME",
        "nome_completo": "Nome Cidadão",
        "numero": 10000 + n,
        "partido_sigla": "PT",
        "cargo": "Deputado Estadual/Distrital",
        "uf": "GO",
        "municipio": None,
        "cpf": "00000000000",
        "cnpj": "",
        "situacao": "AGUARDANDO JULGAMENTO",
        "genero": "MASCULINO",
        "cor_raca": "PARDA",
        "ocupacao": "Médico",
        "grau_instrucao": "SUPERIOR",
        "estado_civil": "CASADO",
        "data_nascimento": "1980-05-12",
        "foto_url": "",
        "coligacao": "COLIGAÇÃO",
        "total_bens_declarados": 1500.75,
        "tem_detalhe": True,
    }


def test_gabarito_vinte_e_cinco_colunas_do_modelo():
    gab = gabarito_modelo()
    assert gab["colunas"] == COLUNAS_MODELO_BASE  # real (arquivo) e fallback alinhados
    assert len(gab["colunas"]) == 25
    assert gab["colunas"][8] == "Membro da FPEvangê\n2023 a 2026\n(Somente para reeleição)"


def test_estrutura_modelo_base_schema():
    candidatos = [_candidato(1), _candidato(2)]
    candidatos[1]["partido_sigla"] = "PL"
    candidatos[1]["nome_urna"] = "BETA"
    df = estrutura_dataframe(candidatos, 2026, "go")
    assert list(df.columns) == gabarito_modelo()["colunas"]
    assert len(df) == 2
    # Ordenação por Partido > Nome.
    assert list(df["Partido"]) == ["PL", "PT"]
    assert list(df["Nome"]) == ["BETA", "NOME"]
    # Inteligência mapeada.
    assert df.loc[1, "Casa Legislativa"] == "Assembleia Legislativa"
    assert df.loc[1, "Eixo de Atuação"] == "Saúde"
    assert df.loc[1, "Idade"] == _idade(_valor_data("1980-05-12"))
    assert df.loc[1, "Observação"].startswith("Bens declarados: R$ 1.500,75")
    assert "Situação do registro: AGUARDANDO JULGAMENTO" in df.loc[1, "Observação"]
    assert df.loc[1, "Sinergia FPE"] == "Neutro"
    assert df.loc[1, "Celular Parlamentar"] == "-"


def test_estrutura_modelo_base_vazia():
    df = estrutura_dataframe([], 2026, "go")
    assert list(df.columns) == gabarito_modelo()["colunas"]
    assert df.empty


def test_bytes_xlsx_aplica_gabarito():
    df = estrutura_dataframe([_candidato(1)], 2026, "GO")
    conteudo = _bytes_xlsx(df)
    assert conteudo[:2] == b"PK"

    import io
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(conteudo))
    sh = wb.active
    assert sh.cell(row=1, column=1).value == "Casa Legislativa"
    assert sh.freeze_panes == "C2"
    assert sh.column_dimensions["B"].width == 36.5
    assert sh.cell(row=1, column=5).value == "Eleição/ Reeleição"


def test_id_eleicao_2026_configuravel():
    assert _id_eleicao(2026) == "2055502026"
    with pytest.raises(ExtrairTSEError):
        _id_eleicao(2019)


def test_validacoes():
    assert _validar_uf(" go ") == "GO"
    with pytest.raises(ExtrairTSEError):
        _validar_uf("XYZ")
    assert _validar_cargo(7, municipais=False) == 7
    with pytest.raises(ExtrairTSEError):
        _validar_cargo(13, municipais=False)


def test_chave_cache_e_derivacoes():
    assert chave_cache(2026, "go", 7) == "tse|2026|GO|7|completo"
    assert chave_cache(2026, "go", 7, limite=50) == "tse|2026|GO|7|limite-50"
    assert _valor_data("1980-05-12") is not None
    assert _valor_data(None) is None
    assert _eixo_atividade("Jornalista e radialista") == "Comunicação e cultura"


def _conteudo_rico():
    return {
        "nomeCompleto": "Fulano de Tal",
        "nomeUrna": "Fulano",
        "cpf": "00000000000",
        "ocupacao": "Médico",
        "grauInstrucao": "SUPERIOR",
        "genero": "MASCULINO",
        "corRaca": "PARDA",
        "dataNascimento": "1980-05-12",
        "estadoCivil": "CASADO",
        "descricaoSituacao": "DEFERIDO",
        "fotoUrl": "http://foto/1.jpg",
        "totalDeBens": 1500.75,
        "bens": [{"descricaoDeTipoDeBem": "CASA", "descricao": "Casa de praia", "valor": 1500.75}],
        "propostas": [{"proposta": "Melhorar a educação"}],
        "redesSociais": [{"titulo": "Instagram", "url": "http://insta/fulano"}],
        "partido": {"sigla": "PSD"},
        "numero": 12345,
        "coligacao": {"nome": "COLIGAÇÃO"},
    }


def test_campos_detalhe_validos_e_maximo():
    assert _normalizar_campos_detalhe(None) is None
    assert _normalizar_campos_detalhe("bens, PROPOSTAS ") == {"bens", "propostas"}


def test_campos_detalhe_invalidos_rejeitam():
    with pytest.raises(ExtrairTSEError):
        _normalizar_campos_detalhe("bens,lixo")


def test_detalhe_rico_maximo_traz_tudo():
    rico = _norm_detalhe_rico(_conteudo_rico(), None)
    assert rico["dados"]["nomeCompleto"] == "Fulano de Tal"
    assert rico["dados"]["situacao"] == "DEFERIDO"
    assert rico["patrimonio"]["totalDeBens"] == 1500.75
    assert rico["patrimonio"]["bens"][0]["valor"] == 1500.75
    assert rico["propostas"][0]["proposta"].startswith("Melhorar")
    assert rico["redesSociais"][0]["url"].startswith("http://insta")
    assert rico["eleicao"]["partido"] == "PSD"


def test_detalhe_rico_restringe_pelos_campos():
    rico = _norm_detalhe_rico(_conteudo_rico(), {"bens"})
    assert "patrimonio" in rico
    assert "dados" not in rico
    assert "propostas" not in rico
    assert "redesSociais" not in rico
    assert rico["eleicao"]["partido"] == "PSD"


def test_bens_individuais_normalizam_valor_e_tipos():
    bens = _bens_individuais({"bens": [{"descricaoDeTipoDeBem": "CASA", "valor": "1000.50"}]})
    assert bens == [{"tipo": "CASA", "descricao": None, "valor": 1000.5}]


def test_detalhe_candidato_round_trip_no_banco():
    database.salvar_detalhe_candidatos(
        CHAVE_GO,
        [{**_candidato(1), "_detalhe_rico": {"dados": {"nomeCompleto": "X"}}}],
    )
    res = database.buscar_detalhe_candidato(CHAVE_GO, "1")
    assert res is not None
    assert res["dados"]["nomeCompleto"] == "X"
    assert database.buscar_detalhe_candidato(CHAVE_GO, "999") is None


def test_cache_atende_do_banco_sem_chamar_api(monkeypatch):
    database.salvar_candidatos_tse(
        CHAVE_GO, 2026, "GO", 7, None, None, [_candidato(1), _candidato(2)]
    )

    async def _falha(*args, **kwargs):
        raise AssertionError("API não deveria ser chamada com cache fresco!")

    monkeypatch.setattr("extrator_tse.extrair_candidatos", _falha)
    dados, origem, chave = asyncio.run(
        obter_candidatos_cacheados(2026, "GO", 7, forcar_atualizacao=False)
    )
    assert origem == "cache"
    assert chave == CHAVE_GO
    assert len(dados) == 2


def test_rota_sincrona_servida_do_cache():
    from fastapi.testclient import TestClient

    import main

    database.salvar_candidatos_tse(
        CHAVE_GO, 2026, "GO", 7, None, None, [_candidato(1), _candidato(2)]
    )
    with TestClient(main.app) as cliente:
        resposta = cliente.get("/tse/exportar/2026/GO/7?sincrono=true")
    assert resposta.status_code == 200
    corpo = resposta.json()
    assert corpo["origem"] == "cache"
    assert corpo["total_candidatos"] == 2
    assert corpo["colunas"] == gabarito_modelo()["colunas"]


def test_trigger_delete_segundo_plano_e_status():
    from fastapi.testclient import TestClient

    import main

    database.salvar_candidatos_tse(
        CHAVE_GO, 2026, "GO", 7, None, None, [_candidato(1), _candidato(2), _candidato(3)]
    )
    with TestClient(main.app) as cliente:
        trigger = cliente.get("/tse/exportar/2026/GO/7")
    assert trigger.status_code == 202
    task_id = trigger.json()["task_id"]
    assert trigger.json()["url_status"] == f"/tse/execucoes/{task_id}"

    registro = None
    for _ in range(100):
        resposta = cliente.get(f"/tse/execucoes/{task_id}")
        assert resposta.status_code == 200
        registro = resposta.json()
        if registro.get("pronto"):
            break
        time.sleep(0.05)

    assert registro is not None
    assert registro["status"] == "Concluído"
    assert registro["total_candidatos"] == 3
    assert len(registro["etapas"]) >= 4


def test_rota_exportar_rejeita_bloco_invalido():
    from fastapi.testclient import TestClient

    import main

    with TestClient(main.app) as cliente:
        resposta = cliente.get("/tse/exportar/2026/GO/7?campos=bens,lixo")
    assert resposta.status_code == 400
    assert "Blocos inválidos" in resposta.json()["detail"]


def test_rota_detalhe_local_404_sem_coleta():
    from fastapi.testclient import TestClient

    import main

    with TestClient(main.app) as cliente:
        resposta = cliente.get(f"/tse/detalhe/{CHAVE_GO}/123")
    assert resposta.status_code == 404


def test_rota_detalhe_local_le_rico_salvo():
    from fastapi.testclient import TestClient

    import main

    database.salvar_detalhe_candidatos(
        CHAVE_GO,
        [{
            **_candidato(1),
            "_detalhe_rico": {
                "dados": {"nomeCompleto": "Fulano de Tal"},
                "patrimonio": {"totalDeBens": 1500.75, "bens": []},
                "eleicao": {"partido": "PSD"},
            },
        }],
    )
    with TestClient(main.app) as cliente:
        resposta = cliente.get(f"/tse/detalhe/{CHAVE_GO}/1")
    assert resposta.status_code == 200
    corpo = resposta.json()
    assert corpo["id_candidato"] == "1"
    assert corpo["detalhe"]["dados"]["nomeCompleto"] == "Fulano de Tal"
    assert corpo["detalhe"]["patrimonio"]["totalDeBens"] == 1500.75


def test_trigger_recusa_execucao_concorrente():
    from fastapi.testclient import TestClient

    import main

    database.criar_execucao_tse("fake-ativo", 2026, "GO", 7)
    database.atualizar_execucao_tse("fake-ativo", status="Executando")

    with TestClient(main.app) as cliente:
        resposta = cliente.get("/tse/exportar/2026/GO/7")
    assert resposta.status_code == 409
    assert resposta.json()["detail"]["execucao"]["task_id"] == "fake-ativo"


def test_status_execucao_desconhecida_404():
    from fastapi.testclient import TestClient

    import main

    with TestClient(main.app) as cliente:
        resposta = cliente.get("/tse/execucoes/task-inexistente")
    assert resposta.status_code == 404


def test_documentos_vazios_seguem_gabarito():
    df = estrutura_dataframe([], 2026, "GO")
    assert isinstance(df, pd.DataFrame)
    assert df.shape == (0, 25)


# ---------------------------------------------------------------------------
# Retry com Exponential Backoff + Jitter
# ---------------------------------------------------------------------------

class _ClienteFake:
    """Cliente HTTP fake: enfileira erros e depois responde 200 com JSON."""

    def __init__(self, erros):
        self.erros = list(erros)
        self.chamadas = 0
        self.atrasos = []

    async def get(self, url, headers=None):
        self.chamadas += 1
        if self.erros:
            problema = self.erros.pop(0)
            if isinstance(problema, int):
                pedido = httpx.Request("GET", url)
                raise httpx.HTTPStatusError(
                    f"erro {problema}", request=pedido, response=httpx.Response(problema)
                )
            raise problema
        return httpx.Response(200, json={"sucesso": True}, request=httpx.Request("GET", url))


async def _dormir_captura(atraso):
    cliente_fake_atual().atrasos.append(atraso)


_cliente_atual: _ClienteFake = _ClienteFake([])


def cliente_fake_atual() -> _ClienteFake:
    return _cliente_atual


def test_atraso_backoff_exponencial_sem_jitter():
    original = (settings.tse_backoff_base, settings.tse_backoff_jitter)
    settings.tse_backoff_base = 2.0
    settings.tse_backoff_jitter = 0.0
    try:
        assert _atraso_tentativa(1) == 2.0
        assert _atraso_tentativa(2) == 4.0
        assert _atraso_tentativa(3) == 8.0
        assert _atraso_tentativa(4) == 16.0
    finally:
        settings.tse_backoff_base, settings.tse_backoff_jitter = original


def test_get_json_rechama_apos_429_e_429(monkeypatch):
    original = (settings.tse_max_tentativas, settings.tse_backoff_base, settings.tse_backoff_jitter)
    settings.tse_max_tentativas = 3
    settings.tse_backoff_base = 0.0
    settings.tse_backoff_jitter = 0.0
    global _cliente_atual
    _cliente_atual = _ClienteFake([429, 429])
    global _dormir_captura
    monkeypatch.setattr("extrator_tse.asyncio.sleep", _dormir_captura)
    try:
        resultado = asyncio.run(_get_json(_cliente_atual, "http://tse.invalido/consultas"))
    finally:
        settings.tse_max_tentativas, settings.tse_backoff_base, settings.tse_backoff_jitter = original
    assert resultado == {"sucesso": True}
    assert _cliente_atual.chamadas == 3
    assert len(_cliente_atual.atrasos) == 2


def test_get_json_desiste_apos_todas_tentativas_403(monkeypatch):
    original = (settings.tse_max_tentativas, settings.tse_backoff_base, settings.tse_backoff_jitter)
    settings.tse_max_tentativas = 3
    settings.tse_backoff_base = 0.0
    settings.tse_backoff_jitter = 0.0
    global _cliente_atual
    _cliente_atual = _ClienteFake([403, 403, 403])
    global _dormir_captura
    monkeypatch.setattr("extrator_tse.asyncio.sleep", _dormir_captura)
    try:
        with pytest.raises(ExtrairTSEError) as exc:
            asyncio.run(_get_json(_cliente_atual, "http://tse.invalido/consultas"))
    finally:
        settings.tse_max_tentativas, settings.tse_backoff_base, settings.tse_backoff_jitter = original
    assert _cliente_atual.chamadas == 3
    assert "403" in str(exc.value)


def test_get_json_rechama_apos_timeout_e_conclui(monkeypatch):
    original = (settings.tse_max_tentativas, settings.tse_backoff_base, settings.tse_backoff_jitter)
    settings.tse_max_tentativas = 3
    settings.tse_backoff_base = 0.0
    settings.tse_backoff_jitter = 0.0
    global _cliente_atual
    _cliente_atual = _ClienteFake([httpx.ConnectTimeout("timeout")])
    global _dormir_captura
    monkeypatch.setattr("extrator_tse.asyncio.sleep", _dormir_captura)
    try:
        resultado = asyncio.run(_get_json(_cliente_atual, "http://tse.invalido/consultas"))
    finally:
        settings.tse_max_tentativas, settings.tse_backoff_base, settings.tse_backoff_jitter = original
    assert resultado == {"sucesso": True}
    assert _cliente_atual.chamadas == 2


def test_get_json_nao_rechama_erro_404_definitivo(monkeypatch):
    original = (settings.tse_max_tentativas, settings.tse_backoff_base, settings.tse_backoff_jitter)
    settings.tse_max_tentativas = 4
    settings.tse_backoff_base = 0.0
    settings.tse_backoff_jitter = 0.0
    global _cliente_atual
    _cliente_atual = _ClienteFake([404])
    global _dormir_captura
    monkeypatch.setattr("extrator_tse.asyncio.sleep", _dormir_captura)
    try:
        with pytest.raises(ExtrairTSEError):
            asyncio.run(_get_json(_cliente_atual, "http://tse.invalido/consultas"))
    finally:
        settings.tse_max_tentativas, settings.tse_backoff_base, settings.tse_backoff_jitter = original
    assert _cliente_atual.chamadas == 1