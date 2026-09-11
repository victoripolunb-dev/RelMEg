"""Testes do Clipping de Novas Proposições + Filtro Inteligente Family Talks."""
import pytest
from docx import Document

from servicos import family_talks
from servicos.exportador_local import (
    ClippingError,
    _caminho_modelo,
    _data_br,
    _data_objeto,
    _dentro_janela,
    _substituir_corpo,
)
from servicos.family_talks import filtrar

CAMARA_ITENS = [
    {
        "siglaTipo": "PL",
        "numero": "595",
        "ano": 2026,
        "id": 42,
        "ementa": "Altera a CLT para ampliar a licença maternidade de 120 para 180 dias.",
        "autor": "Fulano de Tal",
        "data": "2026-09-01T00:00:00",
    },
    {
        "siglaTipo": "PL",
        "numero": "777",
        "ano": 2026,
        "id": 43,
        "ementa": "Dispõe sobre alienação parental e guarda compartilhada.",
        "autor": "Ciclana So e Tal",
        "data": "2026-09-02",
    },
]
SENADO_ITENS = [
    {
        "sigla": "PL",
        "numero": "101",
        "ano": 2026,
        "codigo": 7,
        "ementa": "Institui política de atenção integral à primeira infância.",
        "autor": "Senador Beltrano",
        "data": "2026-09-03",
    },
    {
        "sigla": "PL",
        "numero": "202",
        "ano": 2026,
        "codigo": 8,
        "ementa": "Trata da reforma tributária e dos impostos sobre consumo.",
        "autor": "Senadora Beltrana",
        "data": "2026-09-04",
    },
]


def _doc_com_corpo():
    doc = Document()
    doc.add_paragraph("Plano de Ação", style=doc.styles["Title"])
    doc.add_paragraph("Câmara dos Deputados", style=doc.styles["Title"])
    doc.add_paragraph("PL 595/2024")
    doc.add_paragraph("ementa: parametro")
    doc.add_paragraph("autores: parametro")
    doc.add_paragraph("data: parametro")
    doc.add_paragraph("-- marcador --")
    doc.add_paragraph("Senado Federal", style=doc.styles["Title"])
    doc.add_paragraph("Outro texto fixo")
    return doc


def test_caminho_modelo_ausente():
    try:
        model_path = _caminho_modelo()
    except ClippingError as exc:
        assert "MODELO A SER SEGUIDO" in str(exc)
        pytest.skip("Sem modelo na pasta de entregas — erro orientativo validado")
    assert model_path.name == "MODELO A SER SEGUIDO.docx"


def test_data_br():
    assert _data_br("2026-09-07T00:00:00") == "07/09/2026"
    assert _data_br("2026-09-07") == "07/09/2026"
    assert _data_br(None) is None
    assert _data_br("") is None


def test_data_objeto_normaliza_formatos():
    assert _data_objeto("2026-09-01").isoformat() == "2026-09-01"
    assert _data_objeto("01/09/2026").isoformat() == "2026-09-01"
    assert _data_objeto("2026-09-01T10:30:00").isoformat() == "2026-09-01"
    assert _data_objeto(None) is None
    assert _data_objeto("texto inválido") is None


def test_dentro_janela_inclusiva_e_limites():
    from datetime import date

    ini, fim = date(2026, 8, 31), date(2026, 9, 4)

    # Sem janela: tudo passa.
    assert _dentro_janela("2026-01-01") is True
    assert _dentro_janela(None) is True

    # Dentro da janela (inclusive os extremos).
    assert _dentro_janela("2026-08-31", ini, fim) is True
    assert _dentro_janela("2026-09-04", ini, fim) is True
    assert _dentro_janela("2026-09-01", ini, fim) is True
    assert _dentro_janela("01/09/2026", ini, fim) is True

    # Fora da janela.
    assert _dentro_janela("2026-08-30", ini, fim) is False
    assert _dentro_janela("2026-09-05", ini, fim) is False
    assert _dentro_janela("2025-06-01", ini, fim) is False

    # Data inválida/ausente com janela ativa: não passa.
    assert _dentro_janela(None, ini, fim) is False
    assert _dentro_janela("", ini, fim) is False
    assert _dentro_janela("lixo", ini, fim) is False


def test_janela_inversa_gera_erro(monkeypatch):
    from servicos.exportador_local import _gerar_clipping_async

    async def _falha_busca(keywords: list):
        raise AssertionError("não deveria buscar com janela inválida")

    monkeypatch.setattr("servicos.exportador_local._buscar_camara", _falha_busca)
    monkeypatch.setattr("servicos.exportador_local._buscar_senado", _falha_busca)

    import asyncio

    try:
        asyncio.run(_gerar_clipping_async(
            keywords=["infância"],
            data_inicio=_data_objeto("2026-09-04"),
            data_fim=_data_objeto("2026-08-31"),
        ))
        assert False, "deveria levantar ClippingError"
    except ClippingError as exc:
        assert "data_fim anterior" in str(exc).lower()


def test_gerar_clipping_aplica_janela_de_data(monkeypatch):
    """A janela de datas é aplicada na busca real, antes do Family Talks."""
    from datetime import date

    from servicos.exportador_local import _gerar_clipping_async

    async def _camara_fake(keywords: list):
        return CAMARA_ITENS

    async def _senado_fake(keywords: list):
        # Item FORA da janela: apresentada em 2025, mas com tema válido.
        return SENADO_ITENS + [{
            "sigla": "PL",
            "numero": "999",
            "ano": 2025,
            "codigo": 99,
            "ementa": "Institui política de proteção à primeira infância.",
            "autor": "Senadora Antiga",
            "data": "2025-06-01",
        }]

    def _doc_fake(caminho):
        from unittest.mock import MagicMock
        return MagicMock()

    def _caminho_modelo_fake():
        from pathlib import Path
        return Path("modelo.docx")

    def _pasta_entregas_fake(*args, **kwargs):
        from pathlib import Path
        pasta = Path("C:/entregas")
        pasta.mkdir(parents=True, exist_ok=True)
        return pasta

    monkeypatch.setattr("servicos.exportador_local._buscar_camara", _camara_fake)
    monkeypatch.setattr("servicos.exportador_local._buscar_senado", _senado_fake)
    monkeypatch.setattr("servicos.exportador_local.Document", _doc_fake)
    monkeypatch.setattr("servicos.exportador_local._caminho_modelo", _caminho_modelo_fake)
    monkeypatch.setattr("servicos.exportador_local._pasta_entregas", _pasta_entregas_fake)
    monkeypatch.setattr("servicos.exportador_local._substituir_corpo", lambda *a, **k: None)

    import asyncio

    resultado = asyncio.run(_gerar_clipping_async(
        keywords=["infância"],
        periodo="2026-08-31_a_2026-09-04",
        data_inicio=date(2026, 8, 31),
        data_fim=date(2026, 9, 4),
    ))

    # Nada anterior a 31/08 e nada posterior a 04/09. O item "alienação
    # parental" (CAMARA 2) sai no Family Talks; o de 2025 sai na janela; e o
    # "reforma tributária" (SENADO 4/09) sai por fora de escopo.
    datas = {str(i.get("data"))[:10] for i in resultado["camara"] + resultado["senado"]}
    assert datas == {"2026-09-01", "2026-09-03"}
    assert resultado["total"] == 2
    assert resultado["filtro"]["fora_janela_senado"] == 1
    assert resultado["filtro"]["data_inicio"] == "2026-08-31"
    assert resultado["filtro"]["data_fim"] == "2026-09-04"


def test_filtro_family_talks_escopo():
    aprovados = filtrar(CAMARA_ITENS + SENADO_ITENS)
    ementas = [i["ementa"] for i in aprovados]
    assert len(aprovados) == 2
    assert any("licença maternidade" in e for e in ementas)  # lei: vínculo e parentalidade
    assert any("primeira infância" in e for e in ementas)  # Lei: primeira infância
    assert not any("alienação parental" in e for e in ementas)  # exclusão automática
    assert not any("reforma tributária" in e for e in ementas)  # pauta estranha


def test_substituir_corpo_formata_blocos():
    doc = _doc_com_corpo()
    _substituir_corpo(doc, CAMARA_ITENS, SENADO_ITENS)
    textos = [p.text for p in doc.paragraphs]

    assert "Câmara dos Deputados" in textos
    assert "Senado Federal" in textos
    assert any("PL 595/2026" in t for t in textos)
    assert any("PL 101/2026" in t for t in textos)
    assert any("01/09/2026" in t for t in textos)  # data DD/MM/AAAA da Câmara
    assert any("03/09/2026" in t for t in textos)  # data DD/MM/AAAA do Senado
    assert any("Ciclana So e Tal" in t for t in textos)
    assert not any("-- marcador --" in t for t in textos)
    assert not any("parametro" in t for t in textos)  # placeholders substituídos

    alvos = [
        rel.target_ref
        for rel in doc.part.rels.values()
        if getattr(rel, "is_external", False)
    ]
    assert any("fichadetramitacao?idProposicao=42" in a for a in alvos)
    assert any("/materia/7" in a for a in alvos)


def test_matriz_family_talks_temas_e_exclusoes():
    assert family_talks._TEMAS_PRIORITARIOS  # carregada do backend
    assert family_talks._TEMAS_FORA_ESCOPO  # impedimentos ativos
    aprovados = filtrar(CAMARA_ITENS + SENADO_ITENS)
    assert len(family_talks.temas_detectados(aprovados)) >= 2


def test_rota_clipping_rejeita_template_ausente(monkeypatch):
    from fastapi.testclient import TestClient

    import servicos.exportador_local as exportador_local
    import main

    def _modelo_ausente():
        raise ClippingError(
            "Modelo não encontrado. Confirme que 'MODELO A SER SEGUIDO.docx' "
            "está na pasta backend/templates/ do repositório."
        )

    monkeypatch.setattr(exportador_local, "_caminho_modelo", _modelo_ausente)

    with TestClient(main.app) as cliente:
        resposta = cliente.post(
            "/api/exportar/clipping-semanal",
            params={"keywords": "infância, idosos"},
        )
    assert resposta.status_code == 400
    assert "MODELO A SER SEGUIDO" in resposta.json()["detail"]