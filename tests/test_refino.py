"""Testes de regressão da rodada de refino/auditoria (12/09/2026).

Cobre os fixes: sanitização de nome de arquivo, neutralização de formula
injection na leitura CSV/XLSX, sniff de delimitador, retry seletivo dos
conectores, mapeamento de 404 no hub, dedup-antes-de-enriquecer na Câmara e
inclusão de `ano` no dict TSE. Sem acesso à rede real (APIs mockadas/fakes).
"""
import asyncio
import io

import httpx
import pytest

from routers import planilha as router_planilha
from routers import hub as router_hub
from servicos.exportador_planilha import nome_arquivo_coleta


# ---------------------------------------------------------------------------
# nome_arquivo_coleta — sanitização (achado 5.2)
# ---------------------------------------------------------------------------

def test_nome_arquivo_coleta_sanitiza_separadores():
    nome = nome_arquivo_coleta("../../etc", "SP")
    assert "/" not in nome and "\\" not in nome and ".." not in nome
    assert nome.startswith("Planilha_Coleta_")


def test_nome_arquivo_coleta_preserva_uf_maiuscula():
    assert nome_arquivo_coleta("camara", "sp") == "Planilha_Coleta_camara_SP.xlsx"


def test_nome_arquivo_coleta_defaults():
    assert nome_arquivo_coleta() == "Planilha_Coleta_todas_todas.xlsx"


# ---------------------------------------------------------------------------
# Sniff de delimitador CSV + neutralização de formula injection (F3/6.2)
# ---------------------------------------------------------------------------

def test_ler_csv_sniff_detects_comma():
    conteudo = "nome,uf\nJoao,SP".encode("utf-8")
    registros = router_planilha._ler_csv(conteudo)
    assert registros == [{"nome": "Joao", "uf": "SP"}]


def test_ler_csv_sniff_detects_semicolon():
    conteudo = "nome;uf\nJoao;SP".encode("utf-8")
    registros = router_planilha._ler_csv(conteudo)
    assert registros == [{"nome": "Joao", "uf": "SP"}]


def test_ler_csv_neutraliza_formula_injection():
    conteudo = "nome;valor\nJoao;=SUM(A1:A9)".encode("utf-8")
    registros = router_planilha._ler_csv(conteudo)
    assert registros == [{"nome": "Joao", "valor": "'=SUM(A1:A9)"}]


def test_detectar_delimitador_fallback_ponto_e_virgula():
    assert router_planilha._detectar_delimitador("apenas uma linha de texto") == ";"


def test_ler_xlsx_nao_vaza_formula():
    try:
        import openpyxl
    except ImportError:
        pytest.skip("openpyxl não instalado")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["nome", "valor"])
    ws["A2"] = "Joao"
    ws["B2"] = "=1+1"
    buffer = io.BytesIO()
    wb.save(buffer)
    # data_only=True: fórmulas são lidas pelo valor em cache (sem fórmula bruta),
    # então um = inserido como fórmula nunca vaza o texto injetável ao cliente.
    registros = router_planilha._ler_xlsx(buffer.getvalue())
    assert registros[0]["valor"] != "=1+1"


def test_ler_xlsx_mantem_texto_de_celula_pode_ter_prefixo():
    assert router_planilha._neutralizar_formula("=cmd()") == "'=cmd()"


# ---------------------------------------------------------------------------
# Retry seletivo dos conectores do relmeg_core (E1)
# ---------------------------------------------------------------------------

from relmeg_core.connectors import base_connector as bc  # noqa: E402
from relmeg_core.connectors.camara import CamaraConnector  # noqa: E402


class _FakeClient:
    def __init__(self, respostas):
        self.respostas = list(respostas)
        self.chamadas = 0

    async def get(self, url, params=None, headers=None):
        self.chamadas += 1
        if self.respostas:
            return self.respostas.pop(0)
        return httpx.Response(
            200, json={"ok": True}, request=httpx.Request("GET", url)
        )


class _FakeClienteCM:
    def __init__(self, respostas):
        self.client = _FakeClient(respostas)

    async def __aenter__(self):
        return self.client

    async def __aexit__(self, *args):
        return None


def _resposta(status: int, url: str = "http://fake/api") -> httpx.Response:
    corpo = {"ok": True} if status < 400 else {}
    return httpx.Response(status, json=corpo, request=httpx.Request("GET", url))


def _conector_com_clientes(respostas, max_tentativas: int = 4):
    conn = CamaraConnector()
    conn.max_tentativas = max_tentativas
    conn.backoff_base = 0.0
    conn.backoff_jitter = 0.0
    conexao = _FakeClienteCM(respostas)
    conn._cliente = lambda headers=None: conexao
    return conn, conexao


def test_tentar_nao_retenta_404():
    conn, conexao = _conector_com_clientes([_resposta(404)])
    with pytest.raises(httpx.HTTPStatusError) as exc:
        asyncio.run(conn._tentar(lambda client, url="http://fake": client.get(url)))
    assert exc.value.response.status_code == 404
    assert conexao.client.chamadas == 1


def test_tentar_retenta_429_ate_concluir(monkeypatch):
    conn, conexao = _conector_com_clientes([_resposta(429), _resposta(429), _resposta(200)], max_tentativas=3)
    dormiu = []

    async def _dormir(atraso):
        dormiu.append(atraso)

    monkeypatch.setattr(bc.asyncio, "sleep", _dormir)
    resultado = asyncio.run(conn._tentar(lambda client, url="http://fake": client.get(url)))
    assert resultado == {"ok": True}
    assert conexao.client.chamadas == 3
    assert len(dormiu) == 2


def test_tentar_desiste_apos_tentativas_403(monkeypatch):
    conn, conexao = _conector_com_clientes([_resposta(403), _resposta(403)], max_tentativas=2)

    async def _dormir(atraso):
        pass

    monkeypatch.setattr(bc.asyncio, "sleep", _dormir)
    with pytest.raises(httpx.HTTPStatusError) as exc:
        asyncio.run(conn._tentar(lambda client, url="http://fake": client.get(url)))
    assert exc.value.response.status_code == 403
    assert conexao.client.chamadas == 2


# ---------------------------------------------------------------------------
# Hub: 404 da fonte externa vira HTTP 404 (E1)
# ---------------------------------------------------------------------------

def test_erro_extracao_404_vira_http_404():
    requ = httpx.Request("GET", "http://fake/nao-existe")
    resp = httpx.Response(404, request=requ)
    exc = httpx.HTTPStatusError("não encontrado", request=requ, response=resp)
    erro = router_hub._erro_extracao(exc)
    assert erro.status_code == 404


def test_erro_extracao_502_generico_para_erro_desconhecido():
    erro = router_hub._erro_extracao(RuntimeError("boom"))
    assert erro.status_code == 502


# ---------------------------------------------------------------------------
# exportador_local._buscar_camara — dedup antes de enriquecer (G1)
# ---------------------------------------------------------------------------

def test_buscar_camara_dedup_antes_de_enriquecer(monkeypatch):
    import servicos.exportador_local as exp

    proposicao = {"id": 7, "numero": 42, "ementa": "Marco de energia"}
    chamadas_listagem = [0]
    chamadas_enriquecer = [0]

    async def _lista_fake(siglaTipo=None, ano=None, keywords=None, itens=20, enriquecer=False):
        chamadas_listagem[0] += 1
        return {"proposicoes": [proposicao]}

    async def _enriquece_fake(client, prop):
        chamadas_enriquecer[0] += 1
        prop["autor"] = "Dep. Fulano"
        return prop

    monkeypatch.setattr(exp, "_listar_proposicoes", _lista_fake)
    monkeypatch.setattr(exp, "_enriquecer_camara_autor_data", _enriquece_fake)

    resultado = asyncio.run(exp._buscar_camara(["energia", "tarifa", "energia"]))
    assert len(resultado) == 1
    assert resultado[0]["autor"] == "Dep. Fulano"
    assert chamadas_listagem[0] == 3          # 3 keywords, 3 listagens
    assert chamadas_enriquecer[0] == 1        # apenas 1 proposição única


# ---------------------------------------------------------------------------
# TSE: dict enriquecido carrega o ano da extração (B1)
# ---------------------------------------------------------------------------

class _ClienteTseOk:
    async def get(self, url, headers=None):
        return httpx.Response(
            200,
            json={"candidato": {"ocupacao": "Empresário"}},
            request=httpx.Request("GET", url),
        )


def test_enriquecer_candidato_inclui_ano():
    from servicos.extrator_tse import _enriquecer_candidato

    bruto = {"id": 99, "nomeUrna": "FULANO", "cargo": {"nome": "Deputado Federal"}}
    candidato = asyncio.run(
        _enriquecer_candidato(
            _ClienteTseOk(), 2026, "GO", "2055502026", bruto, "Deputado_Federal"
        )
    )
    assert candidato["ano"] == 2026


# ---------------------------------------------------------------------------
# main.py: documentação protegida por padrão com API key ativa
# ---------------------------------------------------------------------------

def test_caminhos_docs_nao_sao_isentos_por_padrao():
    import main as main_mod

    isentos = main_mod._CAMINHOS_ISENTOS_API_KEY
    assert "/docs" not in isentos
    assert "/openapi.json" not in isentos
    assert "/redoc" not in isentos