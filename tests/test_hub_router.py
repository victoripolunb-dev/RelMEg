"""Testes de integração das rotas do Hub Legislativo (routers/hub.py).

Usa TestClient (FastAPI) com os conectores MOCKADOS — nenhuma chamada a API
externa. Verifica o comportamento on-demand das rotas: disparo via POST,
leitura do repositório via GET, status das fontes e fonte indisponível (501).

O rate limit (slowapi) permanece ligado nas rotas, como exigido pelo AGENTS.md
(rota de disparo = responsável pelas cotas).
"""
import os
import sys
from pathlib import Path

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from slowapi import _rate_limit_exceeded_handler
from slowapi.errors import RateLimitExceeded
from slowapi.middleware import SlowAPIMiddleware

_RAIZ = Path(__file__).resolve().parent.parent
_BACKEND = _RAIZ / "backend"
if str(_BACKEND) not in sys.path:
    sys.path.insert(0, str(_BACKEND))

from rate_limit import limiter
from routers import hub
from relmeg_core.models.schemas import ParlamentarModel, ProjetoDeLeiModel, TramitacaoModel


@pytest.fixture()
def client():
    """App mínimo com as rotas do hub + rate limit real do slowapi."""
    app = FastAPI()
    app.state.limiter = limiter
    app.add_exception_handler(RateLimitExceeded, _rate_limit_exceeded_handler)
    app.add_middleware(SlowAPIMiddleware)
    app.include_router(hub.router)
    with TestClient(app) as c:
        yield c


def _mock_metodo(monkeypatch, fonte: str, nome: str, retorno):
    from relmeg_core.orquestrador import obter_orquestrador

    connector = obter_orquestrador().conector(fonte)

    async def _fake(*args, **kwargs):
        return retorno

    monkeypatch.setattr(connector, nome, _fake)
    return retorno


def _mock_conector(monkeypatch, fonte: str, projeto: dict, tramitacoes: list):
    _mock_metodo(monkeypatch, fonte, "obter_projeto", projeto)
    _mock_metodo(monkeypatch, fonte, "obter_tramitacoes", tramitacoes)
    return projeto, tramitacoes


@pytest.fixture()
def camara_mock(monkeypatch):
    from datetime import date

    projeto = ProjetoDeLeiModel(
        fonte="camara",
        url_origem="https://dadosabertos.camara.leg.br/api/v2/proposicoes/9999",
        id_externo="9999",
        sigla_tipo="PL",
        numero=2571616,
        ano=2025,
        ementa="Dispõe sobre o programa de rastreamento veicular.",
        orgao_origem="camara",
        autor_principal="DEPUTADO FULANO DE TAL",
        pauta_tematica="mobilidade; segurança",
        situacao="Aguardando Designação de Relator(a)",
        comissao_atual="CCJC",
        relator=None,
        data_apresentacao=date(2025, 2, 14),
    )
    tramitacoes = [
        TramitacaoModel(
            fonte="camara",
            url_origem="https://dadosabertos.camara.leg.br/api/v2/proposicoes/9999/tramitacoes",
            id_proposicao_externo="9999",
            data_evento=date(2025, 2, 14),
            orgao_local="Plenário",
            descricao_fase="Apresentação",
            status="Apresentação",
            sequencia=1,
            despacho=None,
        )
    ]
    return _mock_conector(monkeypatch, "camara", projeto, tramitacoes)


@pytest.fixture()
def cldf_mock(monkeypatch):
    from datetime import date

    projeto = ProjetoDeLeiModel(
        fonte="cldf",
        url_origem="https://ple.cl.df.gov.br/pleservico/api/public/proposicao/PL 2473/2026",
        id_externo="PL 2473/2026",
        sigla_tipo="PL",
        numero=2473,
        ano=2026,
        ementa="Ementa de teste da CLDF.",
        orgao_origem="cldf",
        autor_principal="DISTRITAL FULANO",
        situacao="Em tramitação",
        comissao_atual=None,
        relator=None,
    )
    tramitacoes = [
        TramitacaoModel(
            fonte="cldf",
            url_origem="https://ple.cl.df.gov.br/pleservico/api/public/proposicao/PL 2473/2026",
            id_proposicao_externo="PL 2473/2026",
            data_evento=date(2026, 3, 2),
            orgao_local="Plenário",
            descricao_fase="Apresentação",
            status="Apresentação",
            sequencia=1,
            despacho=None,
        )
    ]
    return _mock_conector(monkeypatch, "cldf", projeto, tramitacoes)


def test_status_fontes_sem_rede(client):
    resposta = client.get("/hub/fontes")
    assert resposta.status_code == 200
    dados = resposta.json()
    assert dados["camara"] == "pronta"
    assert dados["algo"].startswith("indisponível")
    assert set(dados) >= {"camara", "senado", "cldf", "algo"}


def test_disparo_sob_demanda_persiste_e_leitura(client, camara_mock):
    projeto, _ = camara_mock

    resposta = client.post("/hub/proposicoes/camara/9999")
    assert resposta.status_code == 201
    corpo = resposta.json()
    assert corpo["fonte"] == "camara"
    assert corpo["id_externo"] == "9999"
    assert corpo["projeto"]["sigla_tipo"] == "PL"
    assert corpo["projeto"]["ementa"] == projeto.ementa
    assert len(corpo["tramitacoes"]) == 1

    leitura = client.get("/hub/proposicoes/camara/9999")
    assert leitura.status_code == 200
    assert leitura.json()["projeto"]["sigla_tipo"] == "PL"
    assert leitura.json()["tramitacoes"][0]["status"] == "Apresentação"


def test_listar_proposicoes_salvas(client, camara_mock):
    client.post("/hub/proposicoes/camara/9999")
    resposta = client.get("/hub/proposicoes/listar?termo=rastreamento")
    assert resposta.status_code == 200
    dados = resposta.json()
    assert dados["total"] >= 1
    assert any(p["sigla_tipo"] == "PL" for p in dados["proposicoes"])


def test_id_com_barra_formato_cldf(client, cldf_mock):
    resposta = client.post("/hub/proposicoes/cldf/PL 2473/2026")
    assert resposta.status_code == 201
    corpo = resposta.json()
    assert corpo["projeto"]["ano"] == 2026

    leitura = client.get("/hub/proposicoes/cldf/PL 2473/2026")
    assert leitura.status_code == 200
    assert leitura.json()["projeto"]["id_externo"] == "PL 2473/2026"


def test_leitura_nao_coletada_retorna_404(client):
    resposta = client.get("/hub/proposicoes/senado/999999")
    assert resposta.status_code == 404


def test_fonte_indisponivel_retorna_501(client):
    resposta = client.post("/hub/proposicoes/algo/PL 1/2025")
    assert resposta.status_code == 501
    detalhe = resposta.json()["detail"].lower()
    assert "algo" in detalhe and "api" in detalhe


def test_fonte_desconhecida_retorna_400(client):
    resposta = client.post("/hub/proposicoes/marte/PL 1/2025")
    assert resposta.status_code == 400


def _parlamentar_fake(fonte: str, id_externo: str, nome: str, uf: str, partido: str = "PSD"):
    return ParlamentarModel(
        fonte=fonte,
        url_origem=f"https://exemplo/{fonte}/{id_externo}",
        id_externo=id_externo,
        nome_completo=nome,
        partido=partido,
        uf=uf,
        nome_urna=nome,
        status_ativo=True,
        cargo="Deputado Federal" if fonte == "camara" else None,
        email="contato@exemplo.br",
    )


def test_parlamentar_sob_demanda_persiste_e_leitura(client, monkeypatch):
    parlamentar = _parlamentar_fake("camara", "111111", "FULANO SILVA", "SP")
    _mock_metodo(monkeypatch, "camara", "obter_parlamentar", parlamentar)

    resposta = client.post("/hub/parlamentares/camara/111111")
    assert resposta.status_code == 201
    corpo = resposta.json()
    assert corpo["parlamentar"]["nome_completo"] == "FULANO SILVA"
    assert corpo["parlamentar"]["uf"] == "SP"

    leitura = client.get("/hub/parlamentares/camara/111111")
    assert leitura.status_code == 200
    assert leitura.json()["parlamentar"]["partido"] == "PSD"


def test_parlamentar_listar_por_filtro(client, monkeypatch):
    _mock_metodo(
        monkeypatch, "camara", "obter_parlamentar",
        _parlamentar_fake("camara", "222222", "MARIA SOUZA", "DF", partido="PL"),
    )
    client.post("/hub/parlamentares/camara/222222")

    resposta = client.get("/hub/parlamentares/listar?uf=DF&partido=PL")
    assert resposta.status_code == 200
    dados = resposta.json()
    assert dados["total"] >= 1
    assert any(p["id_externo"] == "222222" for p in dados["parlamentares"])


def test_parlamentar_nao_coletado_retorna_404(client):
    resposta = client.get("/hub/parlamentares/senado/999999")
    assert resposta.status_code == 404


def test_exportar_ficha_requer_coleta_previa(client):
    resposta = client.post("/hub/exportar/ficha?fonte=senado&id_externo=999999")
    assert resposta.status_code == 404


def test_exportar_ficha_gera_docx_no_padrao_da_casa(client, camara_mock, monkeypatch):
    import os

    caminho_temp = Path(os.environ["DIR_ENTREGAS"]) / "Relatórios"

    client.post("/hub/proposicoes/camara/9999")
    resposta = client.post("/hub/exportar/ficha?fonte=camara&id_externo=9999")
    assert resposta.status_code == 201
    meta = resposta.json()
    arquivo_path = caminho_temp / meta["arquivo"]
    assert meta["arquivo"].startswith("Ficha_Legislativa_PL_2571616")
    assert meta["total_tramitacoes"] == 1
    assert "caminho" not in meta
    assert arquivo_path.exists()

    from docx import Document

    doc = Document(arquivo_path)
    texto = "\n".join(p.text for p in doc.paragraphs)
    assert "FICHA LEGISLATIVA" in texto
    assert "PL 2571616/2025" in texto
    assert "Dispõe sobre o programa de rastreamento veicular." in texto
    assert "Tramitação" in texto
    assert "Identificação" in texto
    assert "RelMeg — extração legislativa sob demanda" in texto


def test_exportar_ficha_parlamentar(client, monkeypatch):
    import os

    _mock_metodo(
        monkeypatch, "camara", "obter_parlamentar",
        _parlamentar_fake("camara", "111111", "FULANO SILVA", "SP"),
    )
    client.post("/hub/parlamentares/camara/111111")

    resposta = client.post("/hub/exportar/ficha-parlamentar?fonte=camara&id_externo=111111")
    assert resposta.status_code == 201
    meta = resposta.json()
    pasta = Path(os.environ["DIR_ENTREGAS"]) / "Relatórios"
    arquivo_path = pasta / meta["arquivo"]
    assert meta["arquivo"].startswith("Ficha_Parlamentar_FULANO_SILVA")
    assert "caminho" not in meta
    assert arquivo_path.exists()

    from docx import Document

    doc = Document(arquivo_path)
    texto = "\n".join(p.text for p in doc.paragraphs)
    assert "FICHA DE PARLAMENTAR" in texto
    assert "FULANO SILVA" in texto
    assert "SP" in texto
    assert "Perfil oficial" in texto


def test_exportar_ficha_parlamentar_requer_coleta(client):
    resposta = client.post("/hub/exportar/ficha-parlamentar?fonte=camara&id_externo=123")
    assert resposta.status_code == 404


def test_exportar_planilha_coleta_gera_xlsx_do_modelo_base(client, monkeypatch):
    from config import settings

    _mock_metodo(
        monkeypatch, "camara", "obter_parlamentar",
        _parlamentar_fake("camara", "111111", "FULANO SILVA", "SP", partido="PL"),
    )
    _mock_metodo(
        monkeypatch, "senado", "obter_parlamentar",
        _parlamentar_fake("senado", "999999", "MARIA SOUZA", "DF", partido="PSD"),
    )
    client.post("/hub/parlamentares/camara/111111")
    client.post("/hub/parlamentares/senado/999999")

    resposta = client.post("/hub/exportar/planilha-coleta?fonte=camara")
    assert resposta.status_code == 201
    meta = resposta.json()
    assert meta["arquivo"].startswith("Planilha_Coleta_camara_")
    assert meta["total"] >= 1
    assert "caminho" not in meta

    destino = settings.dir_perfil / meta["arquivo"]
    assert destino.exists() and destino.stat().st_size > 1000

    import openpyxl

    wb = openpyxl.load_workbook(str(destino))
    sh = wb["Candidatos"]
    assert sh.cell(1, 2).value == "Nome"
    assert sh.cell(2, 1).value == "Câmara"
    assert sh.cell(2, 2).value == "FULANO SILVA"


def test_exportar_planilha_coleta_vazia_gera_cabecalho(client):
    resposta = client.post("/hub/exportar/planilha-coleta?uf=ZZ")
    assert resposta.status_code == 201
    meta = resposta.json()
    assert meta["total"] == 0
    assert meta["arquivo"] == "Planilha_Coleta_todas_ZZ.xlsx"


# ---------------------------------------------------------------------------
# Busca de proposições por palavras-chave (B1 — chamada externa sob demanda)
# ---------------------------------------------------------------------------

def test_busca_proposicoes_camara(client, monkeypatch):
    resultados = [
        {
            "fonte": "camara",
            "url_origem": "https://dadosabertos.camara.leg.br/api/v2/proposicoes/9001",
            "id_externo": "9001",
            "sigla_tipo": "PL",
            "numero": 100,
            "ano": 2026,
            "ementa": "Ementa da busca de teste.",
            "data_apresentacao": "2026-02-14",
            "url_documento": None,
        }
    ]
    _mock_metodo(monkeypatch, "camara", "buscar_proposicoes", resultados)

    resposta = client.get("/hub/busca/proposicoes?fonte=camara&termo=veiculos&itens=5")
    assert resposta.status_code == 200
    corpo = resposta.json()
    assert corpo["fonte"] == "camara"
    assert corpo["termo"] == "veiculos"
    assert corpo["total"] == 1
    assert corpo["resultados"][0]["id_externo"] == "9001"


def test_busca_proposicoes_fonte_sem_api_publica_501(client):
    resposta = client.get("/hub/busca/proposicoes?fonte=algo&termo=teste")
    assert resposta.status_code == 501


def test_busca_proposicoes_dou_200(client, monkeypatch):
    resultados = [
        {
            "fonte": "dou",
            "id_externo": "x/2026/08/28/edicao-1/portaria-1",
            "titulo": "Portaria de concessão.",
            "ementa": "Aprova contrato de concessão de energia.",
            "data_publicacao": "2026-08-28",
            "url": "https://in.gov.br/web/dou/-/x",
        }
    ]
    _mock_metodo(monkeypatch, "dou", "buscar_proposicoes", resultados)

    resposta = client.get("/hub/busca/proposicoes?fonte=dou&termo=concessao&ano=2026&itens=5")
    assert resposta.status_code == 200
    corpo = resposta.json()
    assert corpo["fonte"] == "dou"
    assert corpo["total"] == 1
    assert corpo["resultados"][0]["id_externo"].startswith("x/")


def test_busca_proposicoes_fonte_dou_coleta_por_item_501(client):
    resposta = client.post("/hub/proposicoes/dou/qualquer-publicacao")
    assert resposta.status_code == 501
    assert "dou" in resposta.json()["detail"].lower()


def test_busca_proposicoes_fonte_desconhecida_400(client):
    resposta = client.get("/hub/busca/proposicoes?fonte=zzz&termo=teste")
    assert resposta.status_code == 400


def test_busca_proposicoes_termo_curto_422(client):
    resposta = client.get("/hub/busca/proposicoes?fonte=camara&termo=a")
    assert resposta.status_code == 422